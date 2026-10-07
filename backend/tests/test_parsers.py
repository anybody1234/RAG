from pathlib import Path

import docx
import pymupdf
import pytest

from app.core.language import detect_language
from app.ingestion.parsers import UnsupportedFileError, parse_file, parse_pdf_files
from app.ingestion.pipeline import ingest_file


def texts(parsed) -> list[str]:
    return [block.text for block in parsed.blocks]


def make_pdf(path: Path, pages: list[list[str]], image_only_pages: tuple[int, ...] = ()) -> Path:
    """PDF có lớp text; insert_htmlbox dùng font có sẵn của MuPDF nên hiển thị được tiếng Việt."""
    document = pymupdf.open()
    for index, lines in enumerate(pages, start=1):
        page = document.new_page()
        if index in image_only_pages:
            pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False)
            page.insert_image(pymupdf.Rect(50, 50, 300, 300), pixmap=pixmap)
        else:
            html = "".join(f"<p>{line}</p>" for line in lines)
            page.insert_htmlbox(pymupdf.Rect(50, 40, 550, 800), html, css="p { margin: 0; font-size: 11px; }")
    document.save(path)
    return path


def cong_bao_page(number: int, body: list[str]) -> list[str]:
    return [f"CÔNG BÁO/Số 993 + 994/Ngày 26-12-2019 {number + 2}", *body]


LAW_PAGES = [
    cong_bao_page(1, ["QUỐC HỘI", "Chương I", "NHỮNG QUY ĐỊNH CHUNG", "Điều 1. Phạm vi điều chỉnh",
                      "Bộ luật này quy định tiêu chuẩn lao động.",
                      "Ký bởi: Cổng Thông tin điện tử Chính phủ", "Email: thongtinchinhphu@chinhphu.vn",
                      "Cơ quan: Văn phòng Chính phủ", "Thời gian ký: 31.12.2019 14:43:19 +07:00"]),
    cong_bao_page(2, ["Điều 2. Đối tượng áp dụng", "1. Người lao động.", "2. Người sử dụng lao động."]),
    cong_bao_page(3, ["Điều 3. Giải thích từ ngữ", "Người lao động là người làm việc có trả lương."]),
    cong_bao_page(4, ["Điều 4. Hiệu lực", "Bộ luật này có hiệu lực từ ngày 01 tháng 01 năm 2021.",
                      "(Xem tiếp Công báo số 995 + 996)"]),
]


def test_pdf_law_drops_page_furniture_and_keeps_page_numbers(tmp_path):
    parsed = parse_file(make_pdf(tmp_path / "law.pdf", LAW_PAGES))
    assert parsed.structure == "legal" and parsed.page_count == 4 and parsed.needs_ocr_pages == []
    joined = "\n".join(texts(parsed))
    for furniture in ["CÔNG BÁO", "Ký bởi", "Email:", "Thời gian ký", "Xem tiếp"]:
        assert furniture not in joined
    assert [(b.text, b.page) for b in parsed.blocks if b.text.startswith("Điều")] == [
        ("Điều 1. Phạm vi điều chỉnh", 1), ("Điều 2. Đối tượng áp dụng", 2),
        ("Điều 3. Giải thích từ ngữ", 3), ("Điều 4. Hiệu lực", 4),
    ]


def test_pdf_law_drops_2026_signature_block(tmp_path):
    """Công báo năm 2026 header không có số trang cùng dòng, chữ ký số ghi "Ngày ký" thay cho "Thời gian ký"."""
    pages = [
        ["CÔNG BÁO/Số 18/Ngày 18-01-2026", str(number + 2), *body]
        for number, (_, *body) in enumerate(LAW_PAGES, start=1)
    ]
    pages[0][-4:] = [
        "Ký bởi: CÔNG BÁO NƯỚC CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "Cơ quan: VĂN PHÒNG CHÍNH PHỦ",
        "Ngày ký: 2026-01-17 14:12:26 +07:00",
    ]
    joined = "\n".join(texts(parse_file(make_pdf(tmp_path / "law.pdf", pages))))
    for furniture in ["CÔNG BÁO", "Ký bởi", "Cơ quan:", "Ngày ký"]:
        assert furniture not in joined
    assert "Bộ luật này quy định tiêu chuẩn lao động." in joined


def test_pages_continue_across_files_of_one_document(tmp_path):
    first = make_pdf(tmp_path / "part1.pdf", LAW_PAGES[:2])
    second = make_pdf(tmp_path / "part2.pdf", [["(Tiếp theo Công báo số 993 + 994)"] + LAW_PAGES[2][1:], LAW_PAGES[3]])
    parsed = parse_pdf_files([first, second])
    assert parsed.page_count == 4
    assert [b.page for b in parsed.blocks if b.text.startswith("Điều")] == [1, 2, 3, 4]
    assert "Tiếp theo" not in "\n".join(texts(parsed))


def test_page_without_text_layer_is_flagged_for_ocr(tmp_path):
    parsed = parse_file(make_pdf(tmp_path / "scan.pdf", LAW_PAGES[:2] + [[]], image_only_pages=(3,)))
    assert parsed.needs_ocr_pages == [3]


def test_ocr_engine_fills_scanned_pages(tmp_path):
    class FakeOcr:
        def ocr_page(self, page: pymupdf.Page) -> str:
            return f"Văn bản nhận dạng ở trang {page.number + 1}"

    path = make_pdf(tmp_path / "scan.pdf", LAW_PAGES + [[]], image_only_pages=(5,))
    parsed = parse_pdf_files([path], ocr=FakeOcr())
    assert parsed.needs_ocr_pages == []
    assert parsed.blocks[-1].text == "Văn bản nhận dạng ở trang 5" and parsed.blocks[-1].page == 5


def test_non_legal_pdf_uses_markdown_with_headings(tmp_path):
    pages = [["<h1>User Guide</h1>", "This guide explains how to install the application on a laptop."]]
    parsed = parse_file(make_pdf(tmp_path / "guide.pdf", pages))
    assert parsed.structure == "generic"
    assert any(text.startswith("#") and "User Guide" in text for text in texts(parsed))
    assert all(block.page == 1 for block in parsed.blocks)


def test_docx_headings_paragraphs_and_tables(tmp_path):
    document = docx.Document()
    document.add_heading("Giới thiệu", level=1)
    document.add_paragraph("Nội dung   đầu tiên.")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("A", "B"), ("1", "2")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.save(tmp_path / "doc.docx")
    parsed = parse_file(tmp_path / "doc.docx")
    assert parsed.structure == "generic"
    assert texts(parsed) == ["# Giới thiệu", "Nội dung đầu tiên.", "| A | B |", "| 1 | 2 |"]
    assert parsed.blocks[0].page is None


def test_docx_law_is_detected_as_legal(tmp_path):
    document = docx.Document()
    for n in range(1, 4):
        document.add_paragraph(f"Điều {n}. Tiêu đề {n}")
        document.add_paragraph(f"Nội dung của Điều {n}.")
    document.save(tmp_path / "law.docx")
    assert parse_file(tmp_path / "law.docx").structure == "legal"


def test_html_drops_scripts_and_navigation(tmp_path):
    html = """<html><head><title>T</title><style>p {}</style></head><body>
    <nav>Trang chủ</nav><h1>Tiêu đề
    chính</h1><p>Đoạn
    một <b>đậm</b>.</p><table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>
    <script>alert(1)</script></body></html>"""
    (tmp_path / "page.html").write_text(html, encoding="utf-8")
    assert texts(parse_file(tmp_path / "page.html")) == ["# Tiêu đề chính", "Đoạn một đậm.", "| A | B |", "| 1 | 2 |"]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16"])
def test_text_files_are_decoded(tmp_path, encoding):
    (tmp_path / "note.txt").write_text("Ghi chú tiếng Việt\nDòng hai", encoding=encoding)
    assert texts(parse_file(tmp_path / "note.txt")) == ["Ghi chú tiếng Việt", "Dòng hai"]


def test_markdown_keeps_headings(tmp_path):
    (tmp_path / "readme.md").write_text("# Title\n\nSome text.\n\n## Part\nMore.", encoding="utf-8")
    assert texts(parse_file(tmp_path / "readme.md")) == ["# Title", "Some text.", "## Part", "More."]


def test_old_doc_format_is_rejected_with_conversion_hint(tmp_path):
    (tmp_path / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(UnsupportedFileError, match=r"\.docx"):
        parse_file(tmp_path / "old.doc")


@pytest.mark.parametrize("name, content", [("sheet.xlsx", b"PK"), ("broken.pdf", b"not a pdf")])
def test_unsupported_or_broken_files_raise(tmp_path, name, content):
    (tmp_path / name).write_bytes(content)
    with pytest.raises(UnsupportedFileError):
        parse_file(tmp_path / name)


def test_detect_language():
    assert detect_language("Người lao động có quyền đơn phương chấm dứt hợp đồng lao động.") == "vi"
    assert detect_language("The employee may unilaterally terminate the employment contract.") == "en"


def test_ingest_uploaded_file(tmp_path):
    (tmp_path / "notes.md").write_text("# Ghi chú\nNgười lao động được nghỉ 12 ngày mỗi năm.", encoding="utf-8")
    result = ingest_file(tmp_path / "notes.md", doc_id="u1-notes")
    [chunk] = result.chunks
    assert (chunk.doc_id, chunk.title, chunk.language, chunk.heading_path) == ("u1-notes", "notes", "vi", "Ghi chú")
    assert chunk.page is None and chunk.text.startswith("Ghi chú\n")
