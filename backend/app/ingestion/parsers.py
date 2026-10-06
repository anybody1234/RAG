"""Đọc file upload (PDF, DOCX, HTML, MD, TXT) thành các dòng text kèm số trang.

DOCX và HTML được đưa về dạng markdown đơn giản: heading thành `#`, hàng bảng thành `| ... |`, để chunker chung
chia theo heading. Văn bản luật (nhận ra qua tiêu đề Điều/Article) thì chunk theo Điều dù ở định dạng nào.
"""

import re
from collections.abc import Sequence
from pathlib import Path

import docx
from bs4 import BeautifulSoup, NavigableString
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.legal_chunker import looks_like_legal
from app.ingestion.models import Block, FileFormat, ParsedDocument
from app.ingestion.ocr import OcrEngine
from app.ingestion.pdf import extract_pdf_lines, extract_pdf_markdown
from app.ingestion.text_cleaning import clean_vietnamese_text, expand_ligatures, fix_replacement_chars

FORMATS: dict[str, FileFormat] = {
    ".pdf": "pdf", ".docx": "docx", ".html": "html", ".htm": "html", ".md": "md", ".markdown": "md", ".txt": "txt",
}
DOC_NOT_SUPPORTED = (
    "File .doc (Word 97–2003) chưa được hỗ trợ. Hãy mở file bằng Word, chọn File → Save As → "
    "Word Document (.docx), rồi upload lại file .docx."
)


class UnsupportedFileError(ValueError):
    """File không đọc được; message dùng để báo thẳng cho user."""


def parse_file(path: Path, ocr: OcrEngine | None = None) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".doc":
        raise UnsupportedFileError(DOC_NOT_SUPPORTED)
    file_format = FORMATS.get(suffix)
    if file_format is None:
        raise UnsupportedFileError(
            f"Định dạng {suffix or '(không có đuôi)'} chưa được hỗ trợ. Hỗ trợ: PDF, DOCX, HTML, MD, TXT."
        )
    if file_format == "pdf":
        return parse_pdf_files([path], ocr)

    if file_format == "docx":
        lines = _docx_lines(path)
    elif file_format == "html":
        lines = _html_lines(path.read_bytes())
    else:
        lines = _decode(path.read_bytes()).split("\n")
    text = clean_vietnamese_text(fix_replacement_chars(expand_ligatures("\n".join(lines))))
    blocks = [Block(line, None) for line in text.split("\n") if line]
    structure = "legal" if looks_like_legal(blocks) else "generic"
    return ParsedDocument(format=file_format, structure=structure, blocks=blocks)


def parse_pdf_files(paths: Sequence[Path], ocr: OcrEngine | None = None) -> ParsedDocument:
    """Một văn bản gồm một hoặc nhiều file PDF; số trang đánh liên tục theo thứ tự file."""
    try:
        extracted = extract_pdf_lines(paths, ocr)
    except RuntimeError as exc:  # pymupdf báo file hỏng/mã hoá bằng RuntimeError (FileDataError)
        raise UnsupportedFileError(f"Không đọc được file PDF: {exc}") from exc
    if looks_like_legal(extracted.blocks):
        structure, blocks = "legal", extracted.blocks
    else:
        structure, blocks = "generic", extract_pdf_markdown(paths)
    return ParsedDocument(
        format="pdf",
        structure=structure,
        blocks=blocks,
        page_count=extracted.page_count,
        needs_ocr_pages=extracted.needs_ocr_pages,
    )


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1258", errors="replace")  # bảng mã Windows tiếng Việt


def _docx_heading_level(paragraph: Paragraph) -> int:
    name = (paragraph.style.name if paragraph.style is not None else None) or ""
    if name == "Title":
        return 1
    match = re.fullmatch(r"Heading (\d)", name)
    return min(int(match[1]), 6) if match else 0


def _table_rows(cells_by_row: list[list[str]]) -> list[str]:
    return ["| " + " | ".join(cells) + " |" for cells in cells_by_row if any(cells)]


def _docx_lines(path: Path) -> list[str]:
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # python-docx báo file hỏng bằng nhiều loại lỗi khác nhau
        raise UnsupportedFileError(f"Không đọc được file DOCX: {exc}") from exc
    lines: list[str] = []
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if text:
                level = _docx_heading_level(paragraph)
                lines.append(f"{'#' * level} {text}" if level else text)
        elif tag == "tbl":
            rows = []
            for row in Table(element, document).rows:
                cells, seen = [], set()
                for cell in row.cells:  # ô gộp được trả về nhiều lần
                    if id(cell._tc) not in seen:
                        seen.add(id(cell._tc))
                        cells.append(" ".join(cell.text.split()))
                rows.append(cells)
            lines += _table_rows(rows)
    return lines


_HTML_DROP = ["script", "style", "noscript", "template", "svg", "iframe", "nav", "aside", "footer", "form", "head"]
_HTML_BLOCKS = [
    "p", "div", "li", "ul", "ol", "section", "article", "main", "header", "blockquote", "pre", "dd", "dt",
    "table", "figure", "figcaption", "address",
]


def _html_lines(data: bytes) -> list[str]:
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(_HTML_DROP):
        tag.decompose()
    for node in soup.find_all(string=True):  # xuống dòng trong mã nguồn HTML không phải ngắt đoạn
        if re.search(r"[\r\n]", node) and node.find_parent("pre") is None:
            node.replace_with(NavigableString(re.sub(r"\s+", " ", node)))
    for level in range(1, 7):
        for heading in soup.find_all(f"h{level}"):
            heading.replace_with(NavigableString(f"\n{'#' * level} {' '.join(heading.get_text().split())}\n"))
    for row in soup.find_all("tr"):
        cells = [" ".join(cell.get_text().split()) for cell in row.find_all(["td", "th"])]
        row.replace_with(NavigableString("\n" + "\n".join(_table_rows([cells])) + "\n"))
    for br in soup.find_all("br"):
        br.replace_with(NavigableString("\n"))
    for tag in soup.find_all(_HTML_BLOCKS):
        tag.insert_after(NavigableString("\n"))
    return soup.get_text().split("\n")
